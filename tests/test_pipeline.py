from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

from app.pdf_generator import build_pdf_from_docx


class PipelineSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp(prefix="docx-pdf-test-"))
        self.output_dir = self.temp_dir / "output"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_build_pdf_from_docx(self) -> None:
        docx_path = self.temp_dir / "sample-report.docx"
        image_path = self.temp_dir / "chart.png"

        Image.new("RGB", (640, 360), color=(11, 99, 182)).save(image_path)

        document = Document()
        document.add_paragraph("Observatorio da Industria")
        document.add_paragraph("Panorama estrategico do setor industrial")
        document.add_paragraph("Autores: Equipe de Inteligencia")
        document.add_paragraph("Versao: 1.0")
        document.add_heading("Introducao", level=1)
        document.add_paragraph(
            "Este e um texto de teste para validar a geracao de um relatorio institucional em PDF."
        )
        document.add_heading("Resultados", level=2)
        document.add_paragraph("Insight: A demanda industrial apresentou sinais de recuperacao.")
        document.add_picture(str(image_path), width=Inches(4.8))
        document.add_paragraph("Figura 1 - Evolucao do indicador")
        document.add_paragraph("Fonte: Base sintetica de testes")
        table = document.add_table(rows=3, cols=3)
        table.rows[0].cells[0].text = "Indicador"
        table.rows[0].cells[1].text = "2024"
        table.rows[0].cells[2].text = "2025"
        table.rows[1].cells[0].text = "Produção"
        table.rows[1].cells[1].text = "102"
        table.rows[1].cells[2].text = "110"
        table.rows[2].cells[0].text = "Emprego"
        table.rows[2].cells[1].text = "98"
        table.rows[2].cells[2].text = "101"
        document.add_heading("Referencias", level=1)
        document.add_paragraph("FIEA. Base institucional de demonstracao. Maceio: 2026.")
        document.save(docx_path)

        pdf_path = build_pdf_from_docx(docx_path, self.output_dir)

        self.assertTrue(pdf_path.exists())
        self.assertGreater(pdf_path.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
