"""Build a representative sample DOCX exercising all design-system features."""

from __future__ import annotations

import io
import os
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw


HERE = Path(__file__).resolve().parent


def _make_image(path: Path, color=(11, 99, 182), label="Gráfico de Indicadores") -> None:
    img = Image.new("RGB", (1200, 700), color)
    d = ImageDraw.Draw(img)
    for x in range(0, 1200, 60):
        d.line([(x, 0), (x, 700)], fill=(255, 255, 255), width=1)
    for y in range(0, 700, 60):
        d.line([(0, y), (1200, y)], fill=(255, 255, 255), width=1)
    d.rectangle([60, 60, 1140, 640], outline=(255, 255, 255), width=4)
    bars = [180, 320, 240, 410, 280, 520, 360]
    bw = 100
    gap = (1080 - bw * len(bars)) // (len(bars) - 1)
    x = 80
    for h in bars:
        d.rectangle([x, 620 - h, x + bw, 620], fill=(255, 255, 255))
        x += bw + gap
    d.text((80, 80), label, fill=(255, 255, 255))
    img.save(path)


def build(out_path: Path) -> Path:
    doc = Document()

    cp = doc.core_properties
    cp.title = "Panorama Industrial e Inteligência Estratégica"
    cp.subject = "Análise técnica de indicadores econômicos e de desenvolvimento industrial 2025"
    cp.author = "Observatório da Indústria · Sistema FIEA"
    cp.last_modified_by = "Equipe de Inteligência Competitiva"
    cp.keywords = "indústria; inteligência; análise; economia"
    cp.category = "Relatório Técnico"

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("Ficha Técnica", level=1)
    p = doc.add_paragraph()
    p.add_run("Realização: ").bold = True
    p.add_run("Federação das Indústrias do Estado de Alagoas — Sistema FIEA")
    p = doc.add_paragraph()
    p.add_run("Execução: ").bold = True
    p.add_run("Observatório da Indústria — Núcleo de Inteligência Competitiva")
    p = doc.add_paragraph()
    p.add_run("Coordenação: ").bold = True
    p.add_run("Maria Santos · Carlos Andrade")
    p = doc.add_paragraph()
    p.add_run("Autores: ").bold = True
    p.add_run("João Lima, Beatriz Costa, Rafael Souza, Patrícia Ribeiro")
    p = doc.add_paragraph()
    p.add_run("Revisão Técnica: ").bold = True
    p.add_run("Comitê Editorial do Observatório")
    p = doc.add_paragraph()
    p.add_run("Projeto Gráfico: ").bold = True
    p.add_run("Núcleo de Comunicação Institucional FIEA")

    doc.add_heading("Ficha Catalográfica", level=1)
    doc.add_paragraph(
        "Observatório da Indústria. Panorama Industrial e Inteligência Estratégica / "
        "Observatório da Indústria. — Maceió: Sistema FIEA, 2025."
    )
    doc.add_paragraph("128 p. : il. color.")
    doc.add_paragraph("ISBN 978-85-XXXXX-XX-X")
    doc.add_paragraph("CDU 338.45(813.5)")
    doc.add_paragraph("1. Indústria. 2. Análise Econômica. 3. Inteligência Competitiva.")

    doc.add_heading("Apresentação", level=1)
    doc.add_paragraph(
        "Este relatório apresenta uma análise consolidada do desempenho industrial "
        "no período de referência, organizando indicadores estratégicos, tendências "
        "setoriais e oportunidades identificadas pela rede de inteligência do "
        "Sistema FIEA. A publicação é destinada a tomadores de decisão, "
        "pesquisadores e lideranças empresariais interessadas em compreender, com "
        "rigor técnico, o estado atual do parque industrial e seus vetores de "
        "transformação."
    )
    doc.add_paragraph(
        "A metodologia combina análise quantitativa de séries econômicas oficiais, "
        "leitura qualitativa de fontes setoriais e validação técnica realizada por "
        "comitê de especialistas. As recomendações apresentadas observam o rigor "
        "estatístico e o compromisso institucional do Observatório com a produção "
        "de conhecimento aplicado ao desenvolvimento industrial."
    )
    p = doc.add_paragraph()
    p.add_run("Conselho Editorial — Observatório da Indústria").italic = True

    doc.add_heading("Sumário", level=1)
    doc.add_paragraph("As seções a seguir compõem o relatório:")

    doc.add_heading("Introdução", level=1)
    doc.add_paragraph(
        "A indústria brasileira atravessa um ciclo de reorganização estrutural, "
        "marcado pela aceleração tecnológica, pela transição energética e pela "
        "reconfiguração das cadeias globais de valor. No contexto regional, o "
        "Estado de Alagoas apresenta vetores específicos de oportunidade ligados "
        "ao agronegócio, à química e aos serviços industriais avançados."
    )
    doc.add_heading("Objetivos", level=2)
    doc.add_paragraph(
        "Os objetivos deste estudo compreendem mapear indicadores-chave do "
        "desempenho industrial, identificar gargalos de competitividade e propor "
        "recomendações estratégicas alinhadas à agenda institucional do Sistema FIEA."
    )
    doc.add_paragraph("Os objetivos específicos são:", style="List Bullet")
    doc.add_paragraph("Caracterizar a evolução do PIB industrial estadual.", style="List Bullet")
    doc.add_paragraph("Avaliar produtividade e investimento por setor.", style="List Bullet")
    doc.add_paragraph("Identificar oportunidades de inovação e diversificação.", style="List Bullet")
    doc.add_paragraph("Subsidiar políticas públicas de fomento setorial.", style="List Bullet")

    doc.add_heading("Metodologia", level=1)
    doc.add_paragraph(
        "A pesquisa adotou abordagem quanti-qualitativa, articulando levantamento "
        "documental, análise de séries históricas e entrevistas semiestruturadas. "
        "A triangulação metodológica buscou robustez analítica e segurança nas "
        "inferências, combinando múltiplas fontes e critérios de validação cruzada."
    )
    doc.add_heading("Fontes de Dados", level=2)
    doc.add_paragraph(
        "Foram utilizadas bases oficiais (IBGE, CNI, MDIC, RAIS), além de "
        "publicações setoriais e relatórios regulatórios. A coleta abrangeu "
        "indicadores de produção, emprego, exportação, investimento e inovação."
    )
    doc.add_heading("Procedimentos Analíticos", level=2)
    doc.add_paragraph("As principais etapas de análise foram:")
    doc.add_paragraph("Padronização e limpeza das séries econômicas.", style="List Number")
    doc.add_paragraph("Cálculo de indicadores de variação e participação.", style="List Number")
    doc.add_paragraph("Análise comparativa intersetorial e interestadual.", style="List Number")
    doc.add_paragraph("Validação técnica pelo comitê editorial.", style="List Number")

    p = doc.add_paragraph(style="Quote")
    p.add_run(
        "Insight estratégico: a sofisticação tecnológica é o principal vetor de "
        "diferenciação competitiva da indústria regional no horizonte 2025–2030."
    )

    doc.add_heading("Desenvolvimento", level=1)
    doc.add_paragraph(
        "O desempenho industrial é analisado a seguir a partir de três dimensões "
        "complementares: produção e produtividade, emprego e qualificação, e "
        "exportação e inserção internacional. Cada dimensão é apresentada em "
        "tabela analítica e acompanhada de discussão técnica."
    )

    doc.add_heading("Produção e Produtividade", level=2)
    doc.add_paragraph(
        "A produção industrial estadual apresentou variação consistente no "
        "período, com destaque para os segmentos de químicos, alimentos e "
        "produtos de origem mineral, conforme indicadores apresentados na "
        "Tabela 1."
    )

    p = doc.add_paragraph()
    p.add_run("Tabela 1 — Variação da produção industrial por segmento (2023–2025)").bold = True
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Segmento"
    hdr[1].text = "2023"
    hdr[2].text = "2024"
    hdr[3].text = "2025"
    rows = [
        ("Químicos", "+3,8%", "+5,1%", "+6,2%"),
        ("Alimentos e Bebidas", "+2,1%", "+3,4%", "+4,0%"),
        ("Minerais Não-Metálicos", "+1,2%", "+2,7%", "+3,5%"),
        ("Têxteis", "-0,5%", "+0,8%", "+1,4%"),
        ("Metalurgia", "+0,9%", "+1,6%", "+2,1%"),
    ]
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v

    p = doc.add_paragraph("Fonte: IBGE/PIM-PF; elaboração Observatório da Indústria.")
    p.style = "Caption" if "Caption" in [s.name for s in doc.styles] else None

    img_path = HERE / "_sample_chart.png"
    _make_image(img_path)
    doc.add_picture(str(img_path), width=Inches(5.6))
    doc.add_paragraph(
        "Figura 1 — Evolução do índice de produção industrial estadual."
    )
    doc.add_paragraph(
        "Fonte: Observatório da Indústria · Sistema FIEA."
    )

    doc.add_heading("Emprego e Qualificação", level=2)
    doc.add_paragraph(
        "O emprego industrial formal apresentou crescimento, com variação positiva "
        "concentrada nas ocupações técnicas de nível médio e superior. A "
        "qualificação tornou-se prerrequisito estratégico para a expansão da "
        "produtividade."
    )

    doc.add_heading("Exportação e Inserção Internacional", level=2)
    doc.add_paragraph(
        "As exportações industriais mantiveram trajetória ascendente, com "
        "diversificação parcial da pauta. A inserção em cadeias globais demanda "
        "ações coordenadas de promoção comercial e adequação regulatória."
    )

    p = doc.add_paragraph(style="Quote")
    p.add_run(
        "Recomendação: priorizar agendas de internacionalização articuladas a "
        "programas de qualificação técnica avançada."
    )

    doc.add_heading("Conclusões", level=1)
    doc.add_paragraph(
        "O panorama analisado evidencia oportunidades concretas de adensamento "
        "produtivo e elevação da competitividade industrial. As recomendações "
        "estratégicas estruturam-se em três eixos: produtividade, qualificação e "
        "internacionalização."
    )
    doc.add_paragraph(
        "A continuidade do monitoramento é condição necessária à efetividade das "
        "políticas propostas, demandando articulação institucional permanente "
        "entre o Sistema FIEA, instâncias governamentais e o setor produtivo."
    )

    doc.add_heading("Referências", level=1)
    doc.add_paragraph(
        "CONFEDERAÇÃO NACIONAL DA INDÚSTRIA. Sondagem industrial: panorama 2025. "
        "Brasília: CNI, 2025."
    )
    doc.add_paragraph(
        "INSTITUTO BRASILEIRO DE GEOGRAFIA E ESTATÍSTICA. Pesquisa Industrial "
        "Mensal — Produção Física. Rio de Janeiro: IBGE, 2025."
    )
    doc.add_paragraph(
        "MINISTÉRIO DO DESENVOLVIMENTO, INDÚSTRIA, COMÉRCIO E SERVIÇOS. Boletim "
        "de Comércio Exterior. Brasília: MDIC, 2025."
    )
    doc.add_paragraph(
        "OBSERVATÓRIO DA INDÚSTRIA. Indicadores Estratégicos do Setor Industrial. "
        "Maceió: Sistema FIEA, 2024."
    )

    doc.add_heading("Anexos", level=1)
    p = doc.add_paragraph()
    p.add_run("Anexo A — Glossário de indicadores").bold = True
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Sigla"
    hdr[1].text = "Descrição"
    abbr = [
        ("PIM-PF", "Pesquisa Industrial Mensal — Produção Física (IBGE)"),
        ("RAIS", "Relação Anual de Informações Sociais"),
        ("PIB", "Produto Interno Bruto"),
        ("CNI", "Confederação Nacional da Indústria"),
    ]
    for s, d in abbr:
        cells = table.add_row().cells
        cells[0].text = s
        cells[1].text = d

    doc.add_heading("Contato", level=1)
    doc.add_paragraph("Observatório da Indústria · Sistema FIEA")
    doc.add_paragraph("Avenida Fernandes Lima, S/N — Farol — Maceió/AL")
    doc.add_paragraph("contato@observatorio.fiea.org.br")
    doc.add_paragraph("www.observatorio.fiea.org.br")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


if __name__ == "__main__":
    out = HERE / "sample.docx"
    build(out)
    print(f"wrote {out}")
