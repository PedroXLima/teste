"""
DOCX parser — extracts structured content from a .docx file while preserving
document order (paragraphs and tables interleaved).
"""

import base64
import re
from typing import Any, Dict, List, Optional

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


# ---------------------------------------------------------------------------
# Style classification helpers
# ---------------------------------------------------------------------------

_HEADING_MAP = {
    'heading 1': 'heading1',
    'heading 2': 'heading2',
    'heading 3': 'heading3',
    'heading 4': 'heading4',
    'heading 5': 'heading5',
    'heading 6': 'heading6',
}

_SPECIAL_MAP = {
    'title': 'title',
    'subtitle': 'subtitle',
    'caption': 'caption',
    'intense quote': 'callout',
    'quote': 'quote',
    'block text': 'quote',
}

_TOC_PREFIXES = ('toc ', 'table of contents', 'contents')
_LIST_PREFIXES = ('list', 'bullet', '• ', '- ', 'list paragraph')


def _classify_style(style_name: str) -> str:
    sl = style_name.lower().strip()

    if sl in _HEADING_MAP:
        return _HEADING_MAP[sl]

    for prefix in _HEADING_MAP:
        if sl.startswith(prefix):
            return _HEADING_MAP[prefix]

    if sl in _SPECIAL_MAP:
        return _SPECIAL_MAP[sl]

    for key, val in _SPECIAL_MAP.items():
        if key in sl:
            return val

    for prefix in _TOC_PREFIXES:
        if sl.startswith(prefix):
            return 'toc_entry'

    for prefix in _LIST_PREFIXES:
        if sl.startswith(prefix) or prefix in sl:
            return 'list_item'

    return 'paragraph'


# ---------------------------------------------------------------------------
# DocxParser
# ---------------------------------------------------------------------------

class DocxParser:
    def __init__(self, filepath: str):
        self.filepath = filepath
        self.doc = Document(filepath)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def parse(self) -> Dict[str, Any]:
        content: Dict[str, Any] = {
            'metadata': self._extract_metadata(),
            'elements': self._extract_elements(),
        }
        self._enrich_metadata(content)
        return content

    # ------------------------------------------------------------------
    # Metadata
    # ------------------------------------------------------------------

    def _extract_metadata(self) -> Dict[str, str]:
        props = self.doc.core_properties
        return {
            'title': (props.title or '').strip(),
            'subtitle': '',
            'author': (props.author or '').strip(),
            'subject': (props.subject or '').strip(),
            'description': (getattr(props, 'description', None) or '').strip(),
            'keywords': (getattr(props, 'keywords', None) or '').strip(),
            'created': str(props.created)[:10] if props.created else '',
            'modified': str(props.modified)[:10] if props.modified else '',
            'version': '',
            'reviewer': '',
            'institution': '',
        }

    def _enrich_metadata(self, content: Dict[str, Any]):
        """Pull title/subtitle/version from the first few elements when missing."""
        elements = content['elements']
        meta = content['metadata']

        for elem in elements[:10]:
            etype = elem.get('type', '')
            text = elem.get('text', '').strip()
            if not text:
                continue

            if etype == 'title' and not meta['title']:
                meta['title'] = text
            elif etype == 'subtitle' and not meta['subtitle']:
                meta['subtitle'] = text
            elif etype == 'heading1' and not meta['title']:
                meta['title'] = text

            # Look for version/date/reviewer patterns
            tl = text.lower()
            if not meta['version'] and re.search(r'vers[aã]o|version', tl):
                meta['version'] = text
            if not meta['reviewer'] and re.search(r'revis[oã]|reviewer|revisor', tl):
                meta['reviewer'] = text

    # ------------------------------------------------------------------
    # Element extraction (preserves paragraph/table order via body XML)
    # ------------------------------------------------------------------

    def _extract_elements(self) -> List[Dict[str, Any]]:
        elements: List[Dict[str, Any]] = []

        # Build id→object maps
        para_map = {id(p._element): p for p in self.doc.paragraphs}
        table_map = {id(t._element): t for t in self.doc.tables}

        body = self.doc.element.body
        for child in body:
            local = etree.QName(child.tag).localname

            if local == 'p':
                para = para_map.get(id(child))
                if para is None:
                    continue
                elem = self._process_paragraph(para)
                if elem:
                    elements.append(elem)

            elif local == 'tbl':
                table = table_map.get(id(child))
                if table is None:
                    continue
                elements.append(self._process_table(table))

        return elements

    # ------------------------------------------------------------------
    # Paragraph processing
    # ------------------------------------------------------------------

    def _process_paragraph(self, para) -> Optional[Dict[str, Any]]:
        style_name = para.style.name if para.style else 'Normal'
        text = para.text  # keep original whitespace for now
        stripped = text.strip()

        images = self._extract_images(para)

        # Purely blank paragraphs with no images are skipped
        if not stripped and not images:
            return None

        elem_type = _classify_style(style_name)
        runs = self._process_runs(para)

        # Detect list items by paragraph format even when style says Normal
        if elem_type == 'paragraph' and para.paragraph_format:
            pf = para.paragraph_format
            if pf.left_indent and pf.left_indent.pt and pf.left_indent.pt > 12:
                # Indented paragraph — check for bullet character
                if stripped and stripped[0] in ('•', '-', '–', '▪', '◆', '○', '·'):
                    elem_type = 'list_item'

        return {
            'type': elem_type,
            'style': style_name,
            'text': stripped,
            'raw_text': text,
            'runs': runs,
            'images': images,
        }

    def _process_runs(self, para) -> List[Dict[str, Any]]:
        runs = []
        for run in para.runs:
            if not run.text:
                continue
            runs.append({
                'text': run.text,
                'bold': bool(run.bold),
                'italic': bool(run.italic),
                'underline': bool(run.underline),
            })
        return runs

    def _extract_images(self, para) -> List[Dict[str, str]]:
        images: List[Dict[str, str]] = []
        try:
            for elem in para._element.iter():
                if not isinstance(elem.tag, str):
                    continue
                local = etree.QName(elem.tag).localname
                if local == 'blip':
                    r_embed = elem.get(qn('r:embed'))
                    if r_embed and r_embed in self.doc.part.related_parts:
                        part = self.doc.part.related_parts[r_embed]
                        data = base64.b64encode(part.blob).decode('utf-8')
                        images.append({
                            'data': data,
                            'content_type': part.content_type,
                        })
        except Exception:
            pass
        return images

    # ------------------------------------------------------------------
    # Table processing
    # ------------------------------------------------------------------

    def _process_table(self, table) -> Dict[str, Any]:
        rows: List[List[Dict[str, Any]]] = []
        for i, row in enumerate(table.rows):
            cells: List[Dict[str, Any]] = []
            seen_tc = set()
            for cell in row.cells:
                tc_id = id(cell._element)
                if tc_id in seen_tc:
                    # Skip merged duplicate cells
                    continue
                seen_tc.add(tc_id)
                cells.append({
                    'text': cell.text.strip(),
                    'is_header': i == 0,
                })
            rows.append(cells)

        return {
            'type': 'table',
            'rows': rows,
        }
