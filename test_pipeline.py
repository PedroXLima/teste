#!/usr/bin/env python3
"""End-to-end test: create a sample DOCX, parse it, generate PDF."""

import os
import sys
import tempfile

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.docx_parser import parse_docx
from app.pdf_generator import generate_pdf


def create_sample_docx(path: str):
    """Create a realistic sample DOCX document for testing."""
    doc = Document()

    doc.core_properties.title = "Panorama da Indústria Alagoana"
    doc.core_properties.author = "João Silva, Maria Santos"
    doc.core_properties.subject = "Análise Estratégica do Setor Industrial"

    doc.add_heading("Panorama da Indústria Alagoana", level=1)
    doc.add_heading("Análise Estratégica do Setor Industrial – 2026", level=2)

    doc.add_paragraph("")

    h = doc.add_heading("Realização", level=2)
    doc.add_paragraph("Sistema FIEA – Federação das Indústrias do Estado de Alagoas")
    doc.add_paragraph("Observatório da Indústria")

    doc.add_heading("Coordenação", level=2)
    doc.add_paragraph("Dr. Carlos Eduardo Mendes")

    doc.add_heading("Autores", level=2)
    doc.add_paragraph("João Silva")
    doc.add_paragraph("Maria Santos")
    doc.add_paragraph("Pedro Oliveira")

    doc.add_heading("Revisão", level=2)
    doc.add_paragraph("Ana Beatriz Costa")

    doc.add_heading("Apresentação", level=1)
    doc.add_paragraph(
        "O presente relatório apresenta uma análise abrangente do panorama industrial "
        "do estado de Alagoas, com foco nos principais indicadores econômicos, "
        "tendências de mercado e oportunidades estratégicas para o desenvolvimento "
        "do setor produtivo regional. Este estudo foi elaborado pelo Observatório "
        "da Indústria do Sistema FIEA, com o objetivo de fornecer subsídios para "
        "a tomada de decisão empresarial e a formulação de políticas públicas."
    )
    doc.add_paragraph(
        "A análise contempla dados atualizados sobre emprego, produção, exportações "
        "e investimentos no setor industrial alagoano, utilizando metodologias "
        "consagradas e fontes oficiais de informação. Esperamos que este material "
        "contribua para o fortalecimento da competitividade industrial e o "
        "desenvolvimento sustentável do estado."
    )

    doc.add_heading("Sumário", level=1)
    items = [
        "1. Introdução .......................... 5",
        "2. Metodologia ......................... 8",
        "3. Panorama Econômico .................. 12",
        "4. Setor Industrial .................... 18",
        "5. Conclusões .......................... 25",
        "6. Referências ......................... 28",
        "7. Anexos .............................. 30",
    ]
    for item in items:
        doc.add_paragraph(item)

    doc.add_heading("Introdução", level=1)
    doc.add_paragraph(
        "O setor industrial desempenha papel fundamental na economia alagoana, "
        "respondendo por aproximadamente 22% do Produto Interno Bruto (PIB) estadual "
        "e gerando cerca de 85 mil empregos diretos. A diversificação da base "
        "industrial e a modernização dos processos produtivos são desafios centrais "
        "para a competitividade do estado no cenário nacional."
    )
    doc.add_paragraph(
        "Este relatório analisa as principais tendências e indicadores do setor "
        "industrial alagoano no período de 2023 a 2025, com projeções para 2026. "
        "A análise está estruturada em quatro eixos temáticos: desempenho econômico, "
        "emprego e renda, comércio exterior e investimentos produtivos."
    )

    doc.add_heading("Contexto Macroeconômico", level=2)
    doc.add_paragraph(
        "O cenário macroeconômico brasileiro apresentou sinais de recuperação "
        "moderada nos últimos trimestres, com o PIB nacional crescendo 2,8% em 2025. "
        "A inflação manteve-se dentro da meta estabelecida pelo Banco Central, "
        "e a taxa de juros seguiu trajetória de gradual redução."
    )

    doc.add_heading("Metodologia", level=1)
    doc.add_paragraph(
        "A metodologia adotada neste estudo combina análise quantitativa de dados "
        "secundários com pesquisa qualitativa junto a lideranças empresariais. "
        "As principais fontes de dados utilizadas incluem:"
    )
    doc.add_paragraph("• IBGE – Instituto Brasileiro de Geografia e Estatística")
    doc.add_paragraph("• RAIS/CAGED – Ministério do Trabalho e Emprego")
    doc.add_paragraph("• MDIC – Ministério do Desenvolvimento, Indústria e Comércio")
    doc.add_paragraph("• FIEA – Federação das Indústrias do Estado de Alagoas")

    doc.add_heading("Indicadores Utilizados", level=2)
    doc.add_paragraph(
        "Os indicadores selecionados para esta análise incluem PIB industrial, "
        "número de estabelecimentos, emprego formal, valor da transformação "
        "industrial (VTI), exportações e investimentos anunciados."
    )

    doc.add_heading("Panorama Econômico", level=1)

    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    headers = ["Indicador", "2023", "2024", "2025"]
    for i, h_text in enumerate(headers):
        table.rows[0].cells[i].text = h_text

    data = [
        ["PIB Industrial (R$ bi)", "18,5", "19,2", "20,1"],
        ["Emprego Formal (mil)", "78,3", "81,7", "85,2"],
        ["Exportações (US$ mi)", "1.250", "1.380", "1.490"],
        ["Nº Estabelecimentos", "12.450", "12.890", "13.210"],
        ["VTI (R$ bi)", "8,7", "9,1", "9,6"],
    ]
    for ri, row_data in enumerate(data):
        for ci, cell_val in enumerate(row_data):
            table.rows[ri + 1].cells[ci].text = cell_val

    doc.add_paragraph("Fonte: IBGE/PIA, RAIS/CAGED. Elaboração: Observatório da Indústria/FIEA.")

    doc.add_paragraph(
        "Os dados evidenciam uma trajetória de crescimento consistente do setor "
        "industrial alagoano, com destaque para o aumento do emprego formal "
        "e das exportações."
    )

    doc.add_heading("Distribuição Setorial", level=2)

    table2 = doc.add_table(rows=7, cols=3)
    table2.style = "Table Grid"
    for i, h_text in enumerate(["Setor", "Participação (%)", "Variação 2024-2025"]):
        table2.rows[0].cells[i].text = h_text
    data2 = [
        ["Alimentos e Bebidas", "35,2%", "+3,1%"],
        ["Química e Petroquímica", "22,8%", "+1,5%"],
        ["Têxtil e Confecções", "12,4%", "+2,8%"],
        ["Construção Civil", "10,1%", "+4,2%"],
        ["Metalurgia", "8,5%", "-0,3%"],
        ["Outros", "11,0%", "+1,9%"],
    ]
    for ri, row_data in enumerate(data2):
        for ci, cell_val in enumerate(row_data):
            table2.rows[ri + 1].cells[ci].text = cell_val

    doc.add_paragraph("Fonte: FIEA/Observatório da Indústria, 2025.")

    doc.add_heading("Conclusões", level=1)
    doc.add_paragraph(
        "A análise dos indicadores demonstra que o setor industrial alagoano "
        "mantém uma trajetória de crescimento sustentado, com perspectivas "
        "positivas para 2026. Os principais pontos de atenção incluem:"
    )
    doc.add_paragraph(
        "1. A necessidade de investimentos em inovação e tecnologia para "
        "aumentar a produtividade industrial."
    )
    doc.add_paragraph(
        "2. O fortalecimento das cadeias produtivas locais e a integração "
        "com mercados internacionais."
    )
    doc.add_paragraph(
        "3. A qualificação profissional como fator crítico para a "
        "competitividade do setor."
    )
    doc.add_paragraph(
        "4. O desenvolvimento de infraestrutura logística como elemento "
        "fundamental para a redução de custos operacionais."
    )

    doc.add_heading("Referências", level=1)
    refs = [
        "IBGE. Pesquisa Industrial Anual – PIA. Rio de Janeiro: IBGE, 2025.",
        "FIEA. Relatório Anual da Indústria Alagoana. Maceió: FIEA, 2025.",
        "MDIC. Balança Comercial Brasileira. Brasília: MDIC, 2025.",
        "BANCO CENTRAL DO BRASIL. Relatório de Inflação. Brasília: BCB, 2025.",
        "CNI. Indicadores Industriais. Brasília: CNI, 2025.",
        "SEBRAE. Perfil das Microempresas e Empresas de Pequeno Porte. Brasília: SEBRAE, 2024.",
    ]
    for ref in refs:
        doc.add_paragraph(ref)

    doc.add_heading("Anexos", level=1)
    doc.add_heading("Anexo A – Dados Complementares", level=2)
    doc.add_paragraph(
        "Esta seção apresenta dados complementares sobre o desempenho industrial "
        "alagoano, incluindo séries históricas e indicadores regionais detalhados."
    )

    table3 = doc.add_table(rows=5, cols=3)
    table3.style = "Table Grid"
    for i, h_text in enumerate(["Município", "Nº Indústrias", "Empregos"]):
        table3.rows[0].cells[i].text = h_text
    data3 = [
        ["Maceió", "4.520", "32.100"],
        ["Arapiraca", "1.890", "12.450"],
        ["Palmeira dos Índios", "780", "5.230"],
        ["Penedo", "650", "4.180"],
    ]
    for ri, row_data in enumerate(data3):
        for ci, cell_val in enumerate(row_data):
            table3.rows[ri + 1].cells[ci].text = cell_val

    doc.add_paragraph("Fonte: RAIS/MTE, 2025.")

    doc.save(path)
    print(f"Sample DOCX created: {path}")


def main():
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "sample_report.docx")
        images_dir = os.path.join(tmp, "images")
        pdf_path = os.path.join(tmp, "output.pdf")

        create_sample_docx(docx_path)

        print("Parsing DOCX...")
        parsed = parse_docx(docx_path, images_dir)

        print(f"Title: {parsed['metadata']['title']}")
        print(f"Subtitle: {parsed['metadata']['subtitle']}")
        print(f"Authors: {parsed['metadata']['authors']}")
        print(f"Blocks: {len(parsed['blocks'])}")

        for b in parsed["blocks"][:15]:
            btype = b["type"]
            text = b.get("text", b.get("title", ""))[:60]
            print(f"  [{btype}] {text}")

        print("\nGenerating PDF...")
        generate_pdf(parsed, pdf_path)

        size = os.path.getsize(pdf_path)
        print(f"PDF generated: {pdf_path} ({size:,} bytes)")

        final_path = "/workspace/test_output.pdf"
        import shutil
        shutil.copy2(pdf_path, final_path)
        print(f"Copied to: {final_path}")


if __name__ == "__main__":
    main()
